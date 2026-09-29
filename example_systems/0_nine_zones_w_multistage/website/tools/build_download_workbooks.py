#!/usr/bin/env python3
"""Build standardized Excel downloads from the generated website dataset."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


FILE_NAMES = {
    "epri": "epri-results.xlsx",
    "kmeans_myopic": "genx-kmeans-myopic-results.xlsx",
    "sampled_myopic": "genx-sampled-weeks-myopic-results.xlsx",
}


def load_data(path: Path):
    text = path.read_text()
    prefix = "window.RESULTS_DATA = "
    if not text.startswith(prefix):
        raise ValueError(f"Unexpected data format in {path}")
    payload = text[len(prefix):].rstrip()
    if payload.endswith(";"):
        payload = payload[:-1]
    return json.loads(payload)


def add_sheet(workbook, name, title, headers, rows):
    sheet = workbook.create_sheet(name)
    sheet["A2"] = title
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color="17365D")
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row_index, row in enumerate(rows, start=5):
        for column, value in enumerate(row, start=1):
            cell = sheet.cell(row=row_index, column=column, value=value)
            cell.font = Font(name="Arial", size=10, color="1F2937")
            if isinstance(value, float):
                cell.number_format = "#,##0.000"
    sheet.freeze_panes = "A5"
    for column, header in enumerate(headers, start=1):
        values = [str(header)] + [str(row[column - 1]) for row in rows[:200] if row[column - 1] is not None]
        sheet.column_dimensions[get_column_letter(column)].width = min(max(len(value) for value in values) + 2, 28)
    return sheet


def build_workbook(data, scenario, output: Path):
    workbook = Workbook()
    workbook.remove(workbook.active)
    years = scenario["years"]
    technologies = data["technologies"]
    costs = data["cost_components"]

    add_sheet(
        workbook, "Summary", f"{scenario['label']} results",
        ["Model year", "System cost ($M, 2025 dollars)", "CO2 emissions (Mt)"],
        [[row["year"], row["cost_total"], row["emissions_total"]] for row in years],
    )
    add_sheet(
        workbook, "Costs", f"{scenario['label']} system cost",
        ["Model year", "Reported total ($M, 2025 dollars)"] + [f"{item} ($M, 2025 dollars)" for item in costs],
        [[row["year"], row["cost_total"]] + [row["cost_breakdown"].get(item, 0) for item in costs] for row in years],
    )
    for sheet_name, key, unit, title in (
        ("Capacity", "installed_capacity", "GW", "total installed capacity"),
        ("Buildout", "capacity_mix", "GW", "cumulative buildout"),
        ("Generation", "generation_mix", "TWh", "annual generation"),
    ):
        add_sheet(
            workbook, sheet_name, f"{scenario['label']} {title}",
            ["Model year"] + [f"{item} ({unit})" for item in technologies],
            [[row["year"]] + [row[key].get(item, 0) for item in technologies] for row in years],
        )
    regions = list(dict.fromkeys(region for row in years for region in row.get("emissions_breakdown", {})))
    add_sheet(
        workbook, "Emissions", f"{scenario['label']} annual CO2 emissions",
        ["Model year", "Total (Mt CO2)"] + [f"{item} (Mt CO2)" for item in regions],
        [[row["year"], row["emissions_total"]] + [row["emissions_breakdown"].get(item, 0) for item in regions] for row in years],
    )
    profiles = []
    week_lookup = {week["index"]: week for week in scenario["weeks"]}
    for year_row in years:
        for profile in year_row.get("profiles", []):
            week = week_lookup.get(profile["week"], {})
            for point in profile.get("points", []):
                profiles.append([
                    year_row["year"], profile["week"], week.get("source_week", profile["week"]),
                    point["hour"], point.get("demand"),
                    *[point["generation"].get(item, 0) for item in technologies],
                ])
    add_sheet(
        workbook, "Profiles", f"{scenario['label']} representative-week generation profiles",
        ["Model year", "Representative week index", "Source week", "Hour", "Demand (GW)"]
        + [f"{item} (GW)" for item in technologies], profiles,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--data", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    data = load_data(args.data)
    for scenario in data["scenarios"]:
        filename = FILE_NAMES.get(scenario["id"], f"{scenario['id']}-results.xlsx")
        build_workbook(data, scenario, args.output_dir / filename)
        print(f"Wrote {args.output_dir / filename}")


if __name__ == "__main__":
    main()
